import os
import re
import json
import html
import sqlite3
import secrets
from datetime import datetime
from functools import wraps
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import requests
from dotenv import load_dotenv
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    send_file,
    abort,
)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "instance"
UPLOAD_DIR = DATA_DIR / "uploads"

DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["SECRET_KEY"] = (
    os.getenv("FLASK_SECRET_KEY") or secrets.token_hex(32)
)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024
app.config["UPLOAD_FOLDER"] = str(UPLOAD_DIR)

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
DB_PATH = DATA_DIR / "careerpilot.db"

SKILL_BANK = [
    "python", "java", "c", "c++", "javascript", "typescript",
    "html", "css", "bootstrap", "flask", "django", "fastapi",
    "react", "node.js", "sql", "mysql", "postgresql", "sqlite",
    "mongodb", "firebase", "git", "github", "rest api",
    "api development", "data structures", "algorithms",
    "object-oriented programming", "oop", "dbms",
    "operating systems", "computer networks", "pandas", "numpy",
    "matplotlib", "power bi", "tableau", "excel", "data analysis",
    "machine learning", "deep learning", "nlp", "generative ai",
    "prompt engineering", "langchain", "scikit-learn",
    "tensorflow", "pytorch", "data visualization", "cybersecurity",
    "linux", "docker", "aws", "azure", "communication",
    "project management", "problem solving", "teamwork", "figma",
    "canva", "ui/ux", "testing", "selenium", "arduino", "iot",
    "php", "laravel",
]


# Database connection management
@contextmanager
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")

    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS analyses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                filename TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT '',
                resume_text TEXT NOT NULL,
                score INTEGER NOT NULL,
                skills_json TEXT NOT NULL,
                suggestions_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(user_id)
                    REFERENCES users(id) ON DELETE CASCADE
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS saved_jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                company TEXT NOT NULL DEFAULT '',
                url TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Saved',
                created_at TEXT NOT NULL,
                FOREIGN KEY(user_id)
                    REFERENCES users(id) ON DELETE CASCADE
            )
        """)


init_db()


# Authentication
def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            flash("Please log in to continue.", "warning")
            return redirect(url_for("login"))

        return fn(*args, **kwargs)

    return wrapper


def current_user():
    if not session.get("user_id"):
        return None

    with get_db() as conn:
        return conn.execute(
            "SELECT id, name, email FROM users WHERE id = ?",
            (session["user_id"],),
        ).fetchone()


# Resume processing
def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def extract_text(path):
    suffix = Path(path).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(path)
        chunks = []

        for page in reader.pages:
            chunks.append(page.extract_text() or "")

        return "\n".join(chunks).strip()

    if suffix == ".docx":
        doc = Document(path)
        paragraphs = [
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs).strip()

    return Path(path).read_text(
        encoding="utf-8",
        errors="replace",
    ).strip()


def find_skills(text):
    lower = re.sub(r"\s+", " ", text.lower())
    found = []

    for skill in SKILL_BANK:
        pattern = (
            r"(?<![a-z0-9+#.])"
            + re.escape(skill)
            + r"(?![a-z0-9+#.])"
        )

        if re.search(pattern, lower):
            found.append(skill)

    return sorted(set(found))


def make_suggestions(text, skills, score):
    suggestions = []
    lower = text.lower()

    if len(text.split()) < 120:
        suggestions.append(
            "Add more relevant detail: projects, internship "
            "responsibilities, measurable outcomes, and technical tools."
        )

    if not any(
        word in lower
        for word in ["project", "developed", "built", "implemented"]
    ):
        suggestions.append(
            "Add a Projects section and describe what you built, "
            "the technology used, and the result."
        )

    if not any(
        word in lower
        for word in ["experience", "internship", "work experience"]
    ):
        suggestions.append(
            "If applicable, add internships, volunteering, "
            "freelance work, or practical training."
        )

    if not any(
        word in lower
        for word in ["education", "degree", "bachelor", "diploma"]
    ):
        suggestions.append(
            "Add an Education section with your degree, institute, "
            "graduation year, and relevant coursework."
        )

    if not any(
        word in lower
        for word in ["github", "linkedin", "portfolio"]
    ):
        suggestions.append(
            "Add clickable GitHub, LinkedIn, or portfolio links "
            "so recruiters can verify your work."
        )

    if not re.search(r"\b\d+%?\b", text):
        suggestions.append(
            "Add honest numbers where possible, such as records "
            "processed, time saved, or project users. Never invent metrics."
        )

    if len(skills) < 5:
        suggestions.append(
            "Include a focused Technical Skills section with tools "
            "you can confidently explain in an interview."
        )

    if not any(
        word in lower for word in ["summary", "objective", "profile"]
    ):
        suggestions.append(
            "Add a 2–3 line professional summary tailored to the role."
        )

    if not suggestions:
        suggestions = [
            "Tailor your summary and top project bullets to each job description.",
            "Keep formatting consistent and verify every link and date.",
            "Use concise action verbs and quantify outcomes only when "
            "the numbers are accurate.",
        ]

    return suggestions[:8]


def calculate_score(text, skills):
    lower = text.lower()

    checks = [
        (bool(re.search(r"\b(summary|objective|profile)\b", lower)), 10),
        (
            bool(
                re.search(
                    r"\b(education|degree|diploma|university|college)\b",
                    lower,
                )
            ),
            15,
        ),
        (
            bool(
                re.search(
                    r"\b(project|developed|built|implemented)\b",
                    lower,
                )
            ),
            15,
        ),
        (
            bool(
                re.search(
                    r"\b(experience|internship|volunteer|work experience)\b",
                    lower,
                )
            ),
            15,
        ),
        (len(skills) >= 5, 15),
        (
            bool(re.search(r"\b(github|linkedin|portfolio)\b", lower)),
            10,
        ),
        (bool(re.search(r"\b\d+%?\b", text)), 5),
        (len(text.split()) >= 120, 10),
        (
            bool(
                re.search(
                    r"\b(certification|certificate|training)\b",
                    lower,
                )
            ),
            5,
        ),
    ]

    score = sum(points for passed, points in checks if passed)
    return min(score, 100)


# Groq AI integration
def groq_chat(
    prompt,
    system_prompt=(
        "You are CareerPilot AI, a practical career assistant. "
        "Be accurate, specific, and concise. Do not invent candidate "
        "experience or credentials."
    ),
):
    api_key = os.getenv("GROQ_API_KEY", "").strip()

    if not api_key:
        return (
            None,
            "GROQ_API_KEY is not configured. Showing built-in offline results.",
        )

    model = os.getenv(
        "GROQ_MODEL",
        "llama-3.3-70b-versatile",
    ).strip()

    try:
        response = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            json={
                "model": model,
                "messages": [
                    {
                        "role": "system",
                        "content": system_prompt,
                    },
                    {
                        "role": "user",
                        "content": prompt,
                    },
                ],
                "temperature": 0.4,
                "max_tokens": 1200,
            },
            timeout=35,
        )

        if response.status_code == 401:
            return (
                None,
                "Groq rejected the API key. Check GROQ_API_KEY in your .env file.",
            )

        if response.status_code == 429:
            return (
                None,
                "Groq free-tier rate limit reached. Try again later; "
                "offline results are available.",
            )

        response.raise_for_status()
        data = response.json()
        content = data["choices"][0]["message"]["content"].strip()

        if not content:
            return (
                None,
                "The AI service returned an empty response. "
                "Showing offline results.",
            )

        return content, None

    except requests.Timeout:
        return (
            None,
            "AI request timed out. Check your internet connection and retry.",
        )

    except requests.RequestException:
        return (
            None,
            "Could not reach the AI service. Check your internet/API status; "
            "offline results are available.",
        )

    except (KeyError, ValueError, TypeError):
        return (
            None,
            "The AI service response was not understood. Showing offline results.",
        )


def offline_interview_questions(role, skills):
    role = role.strip() or "entry-level technology"

    skill_text = (
        ", ".join(skills[:8])
        if skills
        else "programming fundamentals, problem solving, and projects"
    )

    return (
        f"Interview practice for: {role}\n\n"
        f"1. Tell me about yourself and why you are interested in {role}.\n"
        "2. Explain one project you built, your individual contribution, "
        "and the biggest challenge.\n"
        f"3. Which of these skills can you demonstrate with a practical "
        f"example: {skill_text}?\n"
        "4. Describe a bug you faced and how you found the cause.\n"
        "5. How do you learn an unfamiliar tool or framework?\n"
        "6. Explain one core concept related to the role in simple words.\n"
        "7. What would you improve in your strongest project if you had "
        "one more week?\n"
        "8. Describe how you test your work before sharing it.\n\n"
        "Practice tip: answer each question in 60–90 seconds and use real examples."
    )


def safe_json_list(raw):
    try:
        result = json.loads(raw)
        return result if isinstance(result, list) else []
    except (ValueError, TypeError):
        return []


# Template globals
@app.context_processor
def inject_globals():
    return {
        "current_user": current_user(),
        "now_year": datetime.now().year,
    }


# Home
@app.route("/")
def index():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    return render_template("index.html")


# Registration
@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        if len(name) < 2 or len(name) > 80:
            flash("Enter a name between 2 and 80 characters.", "danger")

        elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            flash("Enter a valid email address.", "danger")

        elif len(password) < 8:
            flash("Password must contain at least 8 characters.", "danger")

        else:
            try:
                with get_db() as conn:
                    cur = conn.execute(
                        """
                        INSERT INTO users
                        (name, email, password_hash, created_at)
                        VALUES (?, ?, ?, ?)
                        """,
                        (
                            name,
                            email,
                            generate_password_hash(password),
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )

                    session.clear()
                    session["user_id"] = cur.lastrowid

                flash(
                    "Account created. Welcome to CareerPilot AI!",
                    "success",
                )
                return redirect(url_for("dashboard"))

            except sqlite3.IntegrityError:
                flash(
                    "An account with this email already exists.",
                    "danger",
                )

    return render_template("register.html")


# Login
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        with get_db() as conn:
            user = conn.execute(
                "SELECT * FROM users WHERE email = ?",
                (email,),
            ).fetchone()

        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]

            flash("Welcome back!", "success")
            return redirect(url_for("dashboard"))

        flash("Email or password is incorrect.", "danger")

    return render_template("login.html")


# Logout
@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "info")
    return redirect(url_for("index"))


# Dashboard
@app.route("/dashboard")
@login_required
def dashboard():
    with get_db() as conn:
        analyses = conn.execute(
            """
            SELECT * FROM analyses
            WHERE user_id = ?
            ORDER BY id DESC LIMIT 5
            """,
            (session["user_id"],),
        ).fetchall()

        jobs = conn.execute(
            """
            SELECT * FROM saved_jobs
            WHERE user_id = ?
            ORDER BY id DESC LIMIT 5
            """,
            (session["user_id"],),
        ).fetchall()

        total_analyses = conn.execute(
            "SELECT COUNT(*) FROM analyses WHERE user_id = ?",
            (session["user_id"],),
        ).fetchone()[0]

        total_jobs = conn.execute(
            "SELECT COUNT(*) FROM saved_jobs WHERE user_id = ?",
            (session["user_id"],),
        ).fetchone()[0]

    return render_template(
        "dashboard.html",
        analyses=analyses,
        jobs=jobs,
        total_analyses=total_analyses,
        total_jobs=total_jobs,
    )


# Resume analysis
@app.route("/analyze", methods=["GET", "POST"])
@login_required
def analyze():
    if request.method == "POST":
        uploaded = request.files.get("resume")
        role = request.form.get("role", "").strip()[:120]

        if not uploaded or not uploaded.filename:
            flash("Choose a PDF, DOCX, or TXT resume file.", "danger")
            return redirect(url_for("analyze"))

        if not allowed_file(uploaded.filename):
            flash("Unsupported file type. Use PDF, DOCX, or TXT.", "danger")
            return redirect(url_for("analyze"))

        original_name = secure_filename(uploaded.filename) or "resume"
        stored_name = (
            f"{session['user_id']}_{secrets.token_hex(8)}_{original_name}"
        )
        path = UPLOAD_DIR / stored_name

        try:
            uploaded.save(path)
            resume_text = extract_text(path)

            if len(resume_text.strip()) < 40:
                flash(
                    "We could not extract enough text. For scanned PDFs, "
                    "use a text-based PDF or upload DOCX/TXT.",
                    "danger",
                )
                return redirect(url_for("analyze"))

            resume_text = resume_text[:50000]
            skills = find_skills(resume_text)
            score = calculate_score(resume_text, skills)
            suggestions = make_suggestions(resume_text, skills, score)

            with get_db() as conn:
                cur = conn.execute(
                    """
                    INSERT INTO analyses
                    (
                        user_id, filename, role, resume_text, score,
                        skills_json, suggestions_json, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        session["user_id"],
                        original_name,
                        role,
                        resume_text,
                        score,
                        json.dumps(skills),
                        json.dumps(suggestions),
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )

                analysis_id = cur.lastrowid

            return redirect(
                url_for("analysis_detail", analysis_id=analysis_id)
            )

        except Exception as exc:
            app.logger.exception(
                "Resume processing failed: %s",
                exc,
            )
            flash(
                "Resume processing failed. Ensure the file is not corrupted "
                "and try again.",
                "danger",
            )

        finally:
            try:
                if path.exists():
                    path.unlink()
            except OSError:
                pass

    return render_template("analyze.html")


# Analysis details
@app.route("/analysis/<int:analysis_id>")
@login_required
def analysis_detail(analysis_id):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT * FROM analyses
            WHERE id = ? AND user_id = ?
            """,
            (analysis_id, session["user_id"]),
        ).fetchone()

    if not row:
        abort(404)

    skills = safe_json_list(row["skills_json"])
    suggestions = safe_json_list(row["suggestions_json"])

    return render_template(
        "analysis_detail.html",
        analysis=row,
        skills=skills,
        suggestions=suggestions,
    )


# Job description matching
@app.route("/analysis/<int:analysis_id>/match", methods=["GET", "POST"])
@login_required
def job_match(analysis_id):
    with get_db() as conn:
        analysis = conn.execute(
            """
            SELECT * FROM analyses
            WHERE id = ? AND user_id = ?
            """,
            (analysis_id, session["user_id"]),
        ).fetchone()

    if not analysis:
        abort(404)

    result = None

    if request.method == "POST":
        jd = request.form.get("job_description", "").strip()[:30000]

        if len(jd) < 40:
            flash(
                "Paste a job description with at least 40 characters.",
                "danger",
            )
        else:
            resume_skills = set(
                safe_json_list(analysis["skills_json"])
            )
            jd_skills = set(find_skills(jd))

            matched = sorted(resume_skills & jd_skills)
            missing = sorted(jd_skills - resume_skills)

            denominator = max(len(jd_skills), 1)
            score = round(100 * len(matched) / denominator)

            ai_prompt = (
                "Compare this resume and job description. Give concise, "
                "honest suggestions. Do not claim the candidate has skills "
                "that are absent. Use headings: Strengths, Gaps, Next steps.\n\n"
                f"RESUME (truncated):\n{analysis['resume_text'][:9000]}\n\n"
                f"JOB DESCRIPTION:\n{jd[:9000]}"
            )

            ai_text, ai_error = groq_chat(ai_prompt)

            result = {
                "score": score,
                "matched": matched,
                "missing": missing,
                "ai_text": ai_text,
                "ai_error": ai_error,
            }

    return render_template(
        "job_match.html",
        analysis=analysis,
        result=result,
    )


# Interview preparation
@app.route("/interview", methods=["GET", "POST"])
@login_required
def interview():
    output = None
    notice = None

    if request.method == "POST":
        role = (
            request.form.get("role", "").strip()[:120]
            or "Entry-level software role"
        )

        skills = [
            skill.strip()[:50]
            for skill in request.form.get("skills", "").split(",")
            if skill.strip()
        ][:20]

        prompt = (
            f"Create 10 interview practice questions for the role '{role}'. "
            f"Candidate skills: {', '.join(skills) or 'not specified'}. "
            "Group into Technical, Project, and HR. Add short sample answer "
            "guidance, not fabricated personal claims. Keep beginner-friendly "
            "and concise."
        )

        output, notice = groq_chat(prompt)

        if not output:
            output = offline_interview_questions(role, skills)

    return render_template(
        "interview.html",
        output=output,
        notice=notice,
    )


# Saved jobs tracker
@app.route("/jobs", methods=["GET", "POST"])
@login_required
def jobs():
    if request.method == "POST":
        title = request.form.get("title", "").strip()[:150]
        company = request.form.get("company", "").strip()[:150]
        url = request.form.get("url", "").strip()[:500]
        notes = request.form.get("notes", "").strip()[:1000]

        if not title:
            flash("Job title is required.", "danger")

        elif url and not re.match(r"^https?://", url, re.I):
            flash(
                "Job link must start with http:// or https://.",
                "danger",
            )

        else:
            with get_db() as conn:
                conn.execute(
                    """
                    INSERT INTO saved_jobs
                    (
                        user_id, title, company, url, notes,
                        status, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        session["user_id"],
                        title,
                        company,
                        url,
                        notes,
                        "Saved",
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )

            flash("Job saved to your tracker.", "success")
            return redirect(url_for("jobs"))

    with get_db() as conn:
        all_jobs = conn.execute(
            """
            SELECT * FROM saved_jobs
            WHERE user_id = ?
            ORDER BY id DESC
            """,
            (session["user_id"],),
        ).fetchall()

    return render_template("jobs.html", jobs=all_jobs)


# Update job status
@app.route("/jobs/<int:job_id>/status", methods=["POST"])
@login_required
def update_job_status(job_id):
    status = request.form.get("status", "")

    if status not in {
        "Saved",
        "Applied",
        "Interview",
        "Rejected",
        "Offer",
    }:
        flash("Invalid application status.", "danger")

    else:
        with get_db() as conn:
            cur = conn.execute(
                """
                UPDATE saved_jobs
                SET status = ?
                WHERE id = ? AND user_id = ?
                """,
                (status, job_id, session["user_id"]),
            )
            updated = cur.rowcount

        if updated:
            flash("Application status updated.", "success")
        else:
            abort(404)

    return redirect(url_for("jobs"))


# Delete saved job
@app.route("/jobs/<int:job_id>/delete", methods=["POST"])
@login_required
def delete_job(job_id):
    with get_db() as conn:
        cur = conn.execute(
            """
            DELETE FROM saved_jobs
            WHERE id = ? AND user_id = ?
            """,
            (job_id, session["user_id"]),
        )
        deleted = cur.rowcount

    if not deleted:
        abort(404)

    flash("Job removed.", "info")
    return redirect(url_for("jobs"))


# Download resume report
@app.route("/analysis/<int:analysis_id>/report")
@login_required
def download_report(analysis_id):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT * FROM analyses
            WHERE id = ? AND user_id = ?
            """,
            (analysis_id, session["user_id"]),
        ).fetchone()

    if not row:
        abort(404)

    skills = safe_json_list(row["skills_json"])
    suggestions = safe_json_list(row["suggestions_json"])

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=44,
        leftMargin=44,
        topMargin=44,
        bottomMargin=44,
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="SmallBody",
            parent=styles["BodyText"],
            leading=15,
            spaceAfter=7,
        )
    )

    story = [
        Paragraph(
            "CareerPilot AI — Resume Analysis Report",
            styles["Title"],
        ),
        Spacer(1, 10),
        Paragraph(
            f"<b>File:</b> {html.escape(row['filename'])}",
            styles["SmallBody"],
        ),
        Paragraph(
            f"<b>Target role:</b> "
            f"{html.escape(row['role'] or 'Not specified')}",
            styles["SmallBody"],
        ),
        Paragraph(
            f"<b>Analysis date:</b> {row['created_at'][:10]}",
            styles["SmallBody"],
        ),
        Paragraph(
            f"<b>Resume checklist score:</b> {row['score']}/100",
            styles["Heading2"],
        ),
        Paragraph(
            "This is a rule-based checklist score, not a guaranteed "
            "ATS result or hiring prediction.",
            styles["SmallBody"],
        ),
        Paragraph("Detected skills", styles["Heading2"]),
        Paragraph(
            html.escape(
                ", ".join(skills)
                if skills
                else "No listed skills detected automatically."
            ),
            styles["SmallBody"],
        ),
        Paragraph("Suggestions", styles["Heading2"]),
    ]

    for item in suggestions:
        safe_item = html.escape(str(item))
        story.append(
            Paragraph(
                "• " + safe_item,
                styles["SmallBody"],
            )
        )

    story.extend(
        [
            Spacer(1, 12),
            Paragraph(
                "Privacy note: this report is generated from the resume "
                "text extracted for your account.",
                styles["Italic"],
            ),
        ]
    )

    doc.build(story)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="careerpilot_resume_report.pdf",
        mimetype="application/pdf",
    )


# Error handlers
@app.errorhandler(404)
def not_found(_error):
    return (
        render_template(
            "error.html",
            code=404,
            message="The page or record could not be found.",
        ),
        404,
    )


@app.errorhandler(413)
def too_large(_error):
    flash(
        "File is too large. Maximum upload size is 5 MB.",
        "danger",
    )
    return redirect(request.referrer or url_for("analyze"))


@app.errorhandler(500)
def server_error(_error):
    return (
        render_template(
            "error.html",
            code=500,
            message=(
                "Something went wrong. Please retry. "
                "If it continues, check the terminal logs."
            ),
        ),
        500,
    )


if __name__ == "__main__":
    app.run(
        debug=os.getenv("FLASK_DEBUG", "false").lower() == "true"
    )